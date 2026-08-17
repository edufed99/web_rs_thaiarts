"""Create the admin dashboard report as a styled Excel workbook."""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..schemas.dashboard import DashboardOut


NAVY = "202757"
INDIGO = "625BD6"
TEAL = "25A6A1"
SKY = "4C91D1"
CORAL = "DF6B67"
CREAM = "FFF8E8"
PALE = "EEF1FA"
WHITE = "FFFFFF"
TEXT = "203238"
MUTED = "63717A"
GREEN = "D9EFD8"
YELLOW = "FFF0C2"
RED = "F7D7D5"
THIN_GREY = Side(style="thin", color="D9DFE8")


def build_dashboard_report(payload: DashboardOut, *, generated_by: str = "Admin") -> bytes:
    """Return one auditable ``.xlsx`` snapshot of the dashboard payload."""
    wb = Workbook()
    wb.properties.title = "รายงานศูนย์บริหารข้อมูลและติดตามประสิทธิภาพ AI"
    wb.properties.subject = f"Dashboard report — {payload.range_days} days"
    wb.properties.creator = generated_by or "Admin"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    summary = wb.active
    summary.title = "ภาพรวม"
    _build_summary(summary, payload, generated_by)
    _build_activity_trend(wb.create_sheet("แนวโน้ม"), payload)
    _build_users(wb.create_sheet("ผู้ใช้งาน"), payload)
    _build_search(wb.create_sheet("คำค้น"), payload)
    _build_catalog(wb.create_sheet("หมวดหมู่-บริบท"), payload)
    _build_quality(wb.create_sheet("คุณภาพ"), payload)
    _build_recent(wb.create_sheet("กิจกรรมล่าสุด"), payload)
    _build_readme(wb.create_sheet("คำอธิบาย"), payload, generated_by)

    for ws in wb.worksheets:
        _finalize_sheet(ws)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _build_summary(ws, payload: DashboardOut, generated_by: str) -> None:
    _sheet_title(ws, "รายงานศูนย์บริหารข้อมูลและติดตามประสิทธิภาพ AI", 8)
    ws["A3"] = f"ช่วงข้อมูล {payload.range_days} วันล่าสุด"
    ws["A4"] = "สร้างเมื่อ"
    ws["B4"] = _excel_datetime(payload.generated_at)
    ws["B4"].number_format = "yyyy-mm-dd hh:mm"
    ws["D4"] = "ผู้ส่งออก"
    ws["E4"] = generated_by or "Admin"
    ws["G4"] = "แหล่งข้อมูล"
    ws["H4"] = payload.source

    _section(ws, 6, "ตัวชี้วัดหลัก", 8)
    rows = []
    for tile in (
        payload.kpis.members,
        payload.kpis.performances,
        payload.kpis.indices,
        payload.kpis.points,
        payload.kpis.active_users,
        payload.kpis.sessions,
    ):
        rows.append([
            tile.label,
            tile.raw_value,
            tile.delta_pct / 100 if tile.delta_pct is not None else None,
            _tone_label(tile.tone),
            tile.hint,
        ])
    _write_table(ws, 7, ["ตัวชี้วัด", "ค่า", "เปลี่ยนแปลง", "สถานะ", "คำอธิบาย"], rows, "KpiTable")
    _set_number_format(ws, f"B8:B{7 + len(rows)}", "#,##0")
    _set_number_format(ws, f"C8:C{7 + len(rows)}", "0.0%")

    algo_row = 15
    _section(ws, algo_row, "ประสิทธิภาพเส้นทางค้นหา → รายละเอียด", 8)
    algo = payload.algorithm_kpis
    algo_rows = [
        ["จำนวนการค้นหา", algo.search_total, "ครั้ง"],
        ["ค้นหาแล้วเปิดรายละเอียด", algo.search_to_detail_total, "ครั้ง"],
        ["อัตรา Search → Detail", algo.search_to_detail_pct / 100, "อัตรา"],
        ["จำนวนการเปิดดูรายการ", algo.items_shown_total, "ครั้ง"],
        ["CTR ต่อผลลัพธ์ที่แสดง", algo.ctr_pct / 100, "อัตรา"],
    ]
    _write_table(ws, algo_row + 1, ["ตัวชี้วัด", "ค่า", "หน่วย"], algo_rows, "AlgorithmTable")
    ws.cell(algo_row + 4, 2).number_format = "0.0%"
    ws.cell(algo_row + 6, 2).number_format = "0.0%"

    quality_row = 23
    _section(ws, quality_row, "คุณภาพโมเดลล่าสุด", 8)
    q = payload.model_quality
    quality_rows = [
        ["nDCG@10", q.ndcg10, "คุณภาพลำดับผลลัพธ์"],
        ["HR@10", q.hr10, "สัดส่วนผู้ใช้ที่พบผลลัพธ์ตรงใจ"],
        ["MRR@10", q.mrr10, "อันดับเฉลี่ยของผลลัพธ์ที่เกี่ยวข้องรายการแรก"],
        ["Coverage", q.coverage, "ความครอบคลุมรายการในแค็ตตาล็อก"],
        ["Violation rate", q.violation_rate, "อัตราผลลัพธ์ผิดเงื่อนไขบริบท"],
        ["แหล่งประเมิน", q.source, "online / offline / unavailable"],
        ["จำนวนผู้ใช้ทดสอบ", q.test_user_count, "คน"],
        ["จำนวน interaction ทดสอบ", q.test_interaction_count, "รายการ"],
    ]
    _write_table(ws, quality_row + 1, ["ตัวชี้วัด", "ค่า", "คำอธิบาย"], quality_rows, "SummaryQualityTable")
    _set_number_format(ws, f"B{quality_row + 2}:B{quality_row + 6}", "0.0%")


def _build_activity_trend(ws, payload: DashboardOut) -> None:
    _sheet_title(ws, "แนวโน้มการใช้งานรายวัน", 12)
    trend = payload.trend_30d
    rows = [
        [_excel_date(label), _at(trend.sessions, i), _at(trend.searches, i), _at(trend.ratings, i)]
        for i, label in enumerate(trend.labels)
    ]
    _write_table(ws, 4, ["วันที่", "เซสชัน/กิจกรรม", "การค้นหา", "การให้คะแนน"], rows, "ActivityTrendTable")
    if rows:
        _set_number_format(ws, f"A5:A{4 + len(rows)}", "yyyy-mm-dd")
        chart = LineChart()
        chart.title = "แนวโน้ม Sessions / Searches / Ratings"
        chart.style = 13
        chart.y_axis.title = "จำนวน"
        chart.x_axis.title = "วันที่"
        chart.height = 9
        chart.width = 17
        chart.add_data(Reference(ws, min_col=2, max_col=4, min_row=4, max_row=4 + len(rows)), titles_from_data=True)
        for row_index, label in enumerate(trend.labels, start=5):
            ws.cell(row_index, 5, label)
        ws.column_dimensions["E"].hidden = True
        chart.set_categories(Reference(ws, min_col=5, min_row=5, max_row=4 + len(rows)))
        _set_series_titles(chart, ["เซสชัน/กิจกรรม", "การค้นหา", "การให้คะแนน"])
        ws.add_chart(chart, "F4")


def _build_users(ws, payload: DashboardOut) -> None:
    _sheet_title(ws, "ผู้ใช้งานและช่วงเวลาที่ใช้งาน", 26)
    growth = payload.user_growth
    rows = [
        [_excel_date(label), _at(growth.new_users, i), _at(growth.active_users, i)]
        for i, label in enumerate(growth.labels)
    ]
    _section(ws, 4, "ผู้ใช้ใหม่และผู้ใช้งานประจำวัน", 8)
    _write_table(ws, 5, ["วันที่", "ผู้ใช้ใหม่", "Active users"], rows, "UserGrowthTable")
    if rows:
        _set_number_format(ws, f"A6:A{5 + len(rows)}", "yyyy-mm-dd")
        chart = LineChart()
        chart.title = "ผู้ใช้ใหม่เทียบกับ Active users"
        chart.style = 12
        chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=5, max_row=5 + len(rows)), titles_from_data=True)
        for row_index, label in enumerate(growth.labels, start=6):
            ws.cell(row_index, 4, label)
        ws.column_dimensions["D"].hidden = True
        chart.set_categories(Reference(ws, min_col=4, min_row=6, max_row=5 + len(rows)))
        _set_series_titles(chart, ["ผู้ใช้ใหม่", "Active users"])
        chart.height = 8
        chart.width = 14
        ws.add_chart(chart, "E5")

    heatmap_row = max(22, 8 + len(rows))
    _section(ws, heatmap_row, "Heatmap การใช้งานตามวันและชั่วโมง (UTC)", 26)
    hours = payload.usage_heatmap.hour_labels or [f"{hour:02d}:00" for hour in range(24)]
    ws.cell(heatmap_row + 1, 1, "วัน / เวลา")
    for col, hour in enumerate(hours, start=2):
        ws.cell(heatmap_row + 1, col, hour)
    for day_idx, day_label in enumerate(payload.usage_heatmap.weekday_labels, start=0):
        row = heatmap_row + 2 + day_idx
        ws.cell(row, 1, day_label)
        values = payload.usage_heatmap.matrix[day_idx] if day_idx < len(payload.usage_heatmap.matrix) else []
        for col, value in enumerate(values[:24], start=2):
            ws.cell(row, col, int(value))
    heatmap_range = f"B{heatmap_row + 2}:Y{heatmap_row + 8}"
    ws.conditional_formatting.add(
        heatmap_range,
        CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="B7DED9")),
    )
    _style_header(ws, heatmap_row + 1, 25)


def _build_search(ws, payload: DashboardOut) -> None:
    _sheet_title(ws, "พฤติกรรมการค้นหา", 12)
    _section(ws, 4, "Top 5 keyword ที่ผู้ใช้เลือก", 6)
    keyword_rows = [[row.rank, row.term, row.count] for row in payload.top_keywords.items]
    _write_table(ws, 5, ["อันดับ", "Keyword", "จำนวนครั้ง"], keyword_rows, "TopKeywordTable")
    if keyword_rows:
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Keyword ยอดนิยม"
        chart.style = 10
        chart.height = 7
        chart.width = 12
        chart.add_data(Reference(ws, min_col=3, min_row=5, max_row=5 + len(keyword_rows)), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=2, min_row=6, max_row=5 + len(keyword_rows)))
        chart.x_axis.numFmt = "#,##0"
        chart.x_axis.majorUnit = 1
        chart.x_axis.delete = True
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        _set_series_titles(chart, ["จำนวนครั้ง"])
        chart.legend = None
        ws.add_chart(chart, "E5")

    search_row = max(18, 8 + len(keyword_rows))
    _section(ws, search_row, "คำค้นและสัญญาณต่อเนื่อง", 8)
    term_rows = [
        [row.rank, row.term, row.searches, row.views, row.ratings, row.likes, row.score]
        for row in payload.top_search_terms.items
    ]
    _write_table(
        ws,
        search_row + 1,
        ["อันดับ", "คำค้น", "ค้นหา", "เปิดดู", "ให้คะแนน", "ถูกใจ", "คะแนนถ่วงน้ำหนัก"],
        term_rows,
        "SearchSignalTable",
    )


def _build_catalog(ws, payload: DashboardOut) -> None:
    _sheet_title(ws, "ประเภทชุดการแสดงและบริบทยอดนิยม", 14)
    _section(ws, 4, "ประเภทชุดการแสดง", 6)
    category_rows = [[row.name, row.count, row.pct / 100] for row in payload.popular_categories.items]
    _write_table(ws, 5, ["หมวดหมู่", "จำนวนรายการ", "สัดส่วน"], category_rows, "CategoryTable")
    if category_rows:
        _set_number_format(ws, f"C6:C{5 + len(category_rows)}", "0.0%")
        chart = PieChart()
        chart.title = "สัดส่วนประเภทชุดการแสดง"
        chart.height = 8
        chart.width = 12
        chart.add_data(Reference(ws, min_col=2, min_row=5, max_row=5 + len(category_rows)), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=6, max_row=5 + len(category_rows)))
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        ws.add_chart(chart, "E5")

    context_row = max(20, 8 + len(category_rows))
    _section(ws, context_row, "บริบทย่อยตามจำนวน Recommendation request", 6)
    context_rows = [[row.name, row.count, row.pct / 100] for row in payload.popular_subcontexts.items]
    _write_table(ws, context_row + 1, ["บริบทย่อย", "จำนวน Request", "สัดส่วน"], context_rows, "SubContextTable")
    if context_rows:
        _set_number_format(
            ws,
            f"C{context_row + 2}:C{context_row + 1 + len(context_rows)}",
            "0.0%",
        )
        chart = BarChart()
        chart.type = "bar"
        chart.title = "บริบทย่อยยอดนิยม"
        chart.style = 11
        chart.height = 8
        chart.width = 12
        chart.add_data(
            Reference(ws, min_col=2, min_row=context_row + 1, max_row=context_row + 1 + len(context_rows)),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(ws, min_col=1, min_row=context_row + 2, max_row=context_row + 1 + len(context_rows))
        )
        chart.legend = None
        chart.x_axis.numFmt = "#,##0"
        chart.x_axis.majorUnit = 1
        chart.x_axis.delete = True
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        _set_series_titles(chart, ["จำนวน Request"])
        ws.add_chart(chart, f"E{context_row + 1}")


def _build_quality(ws, payload: DashboardOut) -> None:
    _sheet_title(ws, "คุณภาพโมเดลและคุณภาพข้อมูล", 12)
    q = payload.model_quality
    _section(ws, 4, "คุณภาพโมเดลล่าสุด", 8)
    model_rows = [
        ["nDCG@10", q.ndcg10],
        ["HR@10", q.hr10],
        ["MRR@10", q.mrr10],
        ["Coverage", q.coverage],
        ["Violation rate", q.violation_rate],
    ]
    _write_table(ws, 5, ["Metric", "ค่า"], model_rows, "ModelQualityTable")
    _set_number_format(ws, f"B6:B{5 + len(model_rows)}", "0.0%")
    ws["D5"] = "แหล่งประเมิน"
    ws["E5"] = q.source
    ws["D6"] = "ประเมินเมื่อ"
    ws["E6"] = _excel_datetime(q.ran_at) if q.ran_at else None
    ws["E6"].number_format = "yyyy-mm-dd hh:mm"
    ws["D7"] = "ผู้ใช้ทดสอบ"
    ws["E7"] = q.test_user_count
    ws["D8"] = "Interactions ทดสอบ"
    ws["E8"] = q.test_interaction_count

    trend_row = 13
    _section(ws, trend_row, "แนวโน้มคุณภาพโมเดล 30 วัน", 8)
    trend = payload.quality_trend_30d
    trend_rows = [
        [_excel_date(label), _at(trend.ndcg10, i), _at(trend.hr10, i), _at(trend.mrr10, i)]
        for i, label in enumerate(trend.labels)
    ]
    _write_table(ws, trend_row + 1, ["วันที่", "nDCG@10", "HR@10", "MRR@10"], trend_rows, "QualityTrendTable")
    if trend_rows:
        _set_number_format(
            ws,
            f"A{trend_row + 2}:A{trend_row + 1 + len(trend_rows)}",
            "yyyy-mm-dd",
        )
        _set_number_format(
            ws,
            f"B{trend_row + 2}:D{trend_row + 1 + len(trend_rows)}",
            "0.0%",
        )
        chart = LineChart()
        chart.title = "แนวโน้ม nDCG / HR / MRR"
        chart.style = 13
        chart.y_axis.numFmt = "0%"
        chart.height = 8
        chart.width = 14
        chart.add_data(
            Reference(ws, min_col=2, max_col=4, min_row=trend_row + 1, max_row=trend_row + 1 + len(trend_rows)),
            titles_from_data=True,
        )
        for row_index, label in enumerate(trend.labels, start=trend_row + 2):
            ws.cell(row_index, 5, label)
        ws.column_dimensions["E"].hidden = True
        chart.set_categories(Reference(ws, min_col=5, min_row=trend_row + 2, max_row=trend_row + 1 + len(trend_rows)))
        _set_series_titles(chart, ["nDCG@10", "HR@10", "MRR@10"])
        ws.add_chart(chart, f"F{trend_row + 1}")

    data_row = max(31, trend_row + 4 + len(trend_rows))
    _section(ws, data_row, "คุณภาพและความครบถ้วนของข้อมูล", 8)
    metrics = payload.page_quality.metrics
    headers_row = data_row + 1
    rows = [[row.name, row.value / 100, row.target / 100, None, None] for row in metrics]
    _write_table(ws, headers_row, ["มิติคุณภาพ", "ค่าปัจจุบัน", "เป้าหมาย", "ส่วนต่าง", "ผลประเมิน"], rows, "DataQualityTable")
    for offset, metric in enumerate(metrics, start=1):
        row = headers_row + offset
        ws.cell(row, 4, f"=B{row}-C{row}")
        ws.cell(row, 5, f'=IF(B{row}>=C{row},"ผ่าน","ต้องปรับปรุง")')
        ws.cell(row, 5).fill = PatternFill("solid", fgColor=GREEN if metric.value >= metric.target else RED)
    if metrics:
        _set_number_format(
            ws,
            f"B{headers_row + 1}:D{headers_row + len(metrics)}",
            "0.0%",
        )
    ws.cell(headers_row + len(metrics) + 2, 1, "ประเด็นที่ต้องแก้ไข")
    ws.cell(headers_row + len(metrics) + 2, 2, payload.page_quality.open_issues)


def _build_recent(ws, payload: DashboardOut) -> None:
    _sheet_title(ws, "กิจกรรมล่าสุดในระบบ", 10)
    rows = [
        [row.log_id, _excel_datetime(row.time), row.user, row.action, row.type, row.target]
        for row in payload.recent_activity.items
    ]
    _write_table(ws, 4, ["Log ID", "เวลา", "ผู้ใช้", "Action", "ประเภท", "เป้าหมาย"], rows, "RecentActivityTable")
    if rows:
        _set_number_format(ws, f"B5:B{4 + len(rows)}", "yyyy-mm-dd hh:mm")


def _build_readme(ws, payload: DashboardOut, generated_by: str) -> None:
    _sheet_title(ws, "คำอธิบายรายงานและที่มาของข้อมูล", 8)
    rows = [
        ["ชื่อรายงาน", "ศูนย์บริหารข้อมูลและติดตามประสิทธิภาพ AI"],
        ["ช่วงเวลา", f"{payload.range_days} วันล่าสุด"],
        ["ผู้ส่งออก", generated_by or "Admin"],
        ["แหล่งข้อมูล", payload.source],
        ["API ต้นทาง", f"http://127.0.0.1:8001/metrics/dashboard?range={payload.range_days}d"],
        ["Sessions", "จำนวนกิจกรรมรวมรายวัน; ใช้สำหรับดูแนวโน้มการใช้งาน"],
        ["Active users", "จำนวน user_key ที่ไม่ซ้ำในช่วงเวลา"],
        ["Search → Detail", "ผู้ใช้ที่ค้นหาแล้วเปิดหน้ารายละเอียดในช่วงวิเคราะห์"],
        ["nDCG@10", "คุณภาพการเรียงลำดับผลลัพธ์ 10 อันดับแรก; ค่ายิ่งสูงยิ่งดี"],
        ["HR@10", "สัดส่วนผู้ใช้ที่พบผลลัพธ์เกี่ยวข้องภายใน 10 อันดับแรก"],
        ["MRR@10", "อันดับเฉลี่ยแบบ reciprocal ของผลลัพธ์เกี่ยวข้องรายการแรก"],
        ["Coverage", "สัดส่วนรายการในแค็ตตาล็อกที่ระบบสามารถแนะนำได้"],
        ["Violation rate", "สัดส่วนผลลัพธ์ที่ผิดเงื่อนไขบริบท; ค่ายิ่งต่ำยิ่งดี"],
        ["Top keyword", "คำสำคัญที่ผู้ใช้เลือกใน recommendation request และ search telemetry"],
        ["หมายเหตุ", "รายงานเป็น snapshot ณ เวลาส่งออก ตัวเลขอาจเปลี่ยนเมื่อระบบมีข้อมูลใหม่"],
    ]
    _write_table(ws, 4, ["หัวข้อ", "คำอธิบาย"], rows, "ReadmeTable")
    ws.column_dimensions["B"].width = 90
    for row in ws["B5:B30"]:
        row[0].alignment = Alignment(vertical="top", wrap_text=True)


def _sheet_title(ws, text: str, width: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=width)
    cell = ws.cell(1, 1, text)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name="Leelawadee UI", size=20, bold=True, color=WHITE)
    cell.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 14


def _section(ws, row: int, text: str, width: int) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    cell = ws.cell(row, 1, text)
    cell.fill = PatternFill("solid", fgColor=PALE)
    cell.font = Font(name="Leelawadee UI", size=12, bold=True, color=NAVY)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 24


def _write_table(ws, header_row: int, headers: Sequence[str], rows: Sequence[Sequence[object]], name: str) -> None:
    for col, header in enumerate(headers, start=1):
        ws.cell(header_row, col, header)
    _style_header(ws, header_row, len(headers))
    for row_offset, values in enumerate(rows, start=1):
        for col, value in enumerate(values, start=1):
            ws.cell(header_row + row_offset, col, value)
    if rows:
        ref = f"A{header_row}:{_column_letter(len(headers))}{header_row + len(rows)}"
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)


def _style_header(ws, row: int, col_count: int) -> None:
    for col in range(1, col_count + 1):
        cell = ws.cell(row, col)
        cell.fill = PatternFill("solid", fgColor=INDIGO)
        cell.font = Font(name="Leelawadee UI", bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GREY)
    ws.row_dimensions[row].height = 24


def _finalize_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = None
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            if cell.row > 2 and cell.font == Font():
                cell.font = Font(name="Leelawadee UI", size=10, color=TEXT)
            cell.alignment = Alignment(
                horizontal=cell.alignment.horizontal,
                vertical=cell.alignment.vertical or "center",
                wrap_text=cell.alignment.wrap_text,
            )
    _set_widths(ws)


def _set_widths(ws) -> None:
    caps = {1: 30, 2: 24, 3: 22, 4: 22, 5: 36, 6: 38, 7: 22, 8: 22}
    for col_idx in range(1, min(ws.max_column, 25) + 1):
        column_letter = _column_letter(col_idx)
        existing_width = ws.column_dimensions[column_letter].width
        if existing_width and existing_width > 30:
            continue
        max_len = 0
        for row_idx in range(3, min(ws.max_row, 200) + 1):
            value = ws.cell(row_idx, col_idx).value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        width = min(max(11, max_len + 3), caps.get(col_idx, 14))
        ws.column_dimensions[column_letter].width = width


def _set_number_format(ws, cell_range: str, number_format: str) -> None:
    for row in ws[cell_range]:
        for cell in row:
            cell.number_format = number_format


def _set_series_titles(chart, titles: Iterable[str]) -> None:
    for series, title in zip(chart.series, titles):
        series.tx = SeriesLabel(v=title)


def _tone_label(tone: str) -> str:
    return {"positive": "ดี", "warning": "เฝ้าระวัง", "danger": "ต้องแก้ไข"}.get(tone, "ปกติ")


def _excel_datetime(value: str) -> datetime | str | None:
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        return parsed.replace(tzinfo=None)
    except ValueError:
        return value


def _excel_date(value: str) -> date | str:
    try:
        return date.fromisoformat(value[:10])
    except (TypeError, ValueError):
        return value


def _at(values: Sequence[int | float], index: int) -> int | float:
    return values[index] if index < len(values) else 0


def _column_letter(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result
